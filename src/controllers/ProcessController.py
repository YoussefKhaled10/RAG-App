import csv
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, List
from uuid import UUID

from langchain_community.document_loaders import PyMuPDFLoader, TextLoader

from models import ProcessingEnum
from .BaseController import BaseController
from .ProjectController import ProjectController


@dataclass
class Document:
    page_content: str
    metadata: dict[str, Any]


class ProcessController(BaseController):
    def __init__(
        self,
        tenant_id: UUID,
        project_id: int,
    ):
        super().__init__()
        self.tenant_id = tenant_id
        self.project_id = project_id
        self.project_path = ProjectController().get_project_path(
            tenant_id=tenant_id,
            project_id=project_id,
        )

    def get_file_extension(self, file_id: str) -> str:
        safe_file_id = Path(file_id).name
        return Path(safe_file_id).suffix.lower()

    def get_file_path(self, file_id: str) -> str:
        if not file_id or not file_id.strip():
            raise ValueError("file_id cannot be empty")
        safe_file_id = Path(file_id).name
        if safe_file_id != file_id:
            raise ValueError("invalid file_id")

        project_directory = Path(self.project_path).resolve()
        file_path = (project_directory / safe_file_id).resolve()
        if file_path.parent != project_directory:
            raise ValueError("invalid file path")
        return str(file_path)

    def clean_text(self, text: str | None) -> str:
        if text is None:
            return ""
        cleaned_text = str(text).replace("\x00", "")
        cleaned_text = re.sub(
            r"[\x01-\x08\x0B\x0C\x0E-\x1F\x7F]",
            "",
            cleaned_text,
        )
        cleaned_text = re.sub(r"[ \t]+", " ", cleaned_text)
        cleaned_text = re.sub(r"\n{3,}", "\n\n", cleaned_text)
        return cleaned_text.strip()

    def clean_metadata(
        self,
        metadata: dict[str, Any] | None,
    ) -> dict[str, Any]:
        if not isinstance(metadata, dict):
            return {}
        cleaned_metadata: dict[str, Any] = {}
        for key, value in metadata.items():
            clean_key = self.clean_text(str(key))
            if not clean_key:
                continue
            if isinstance(value, str):
                cleaned_metadata[clean_key] = self.clean_text(value)
            elif value is None:
                cleaned_metadata[clean_key] = ""
            elif isinstance(value, (bool, int, float, list, dict)):
                cleaned_metadata[clean_key] = value
            else:
                cleaned_metadata[clean_key] = str(value)
        return cleaned_metadata

    def _base_metadata(
        self,
        file_id: str,
        file_type: str,
    ) -> dict[str, Any]:
        return {
            "file_id": file_id,
            "stored_file_name": file_id,
            "file_type": file_type,
            "tenant_id": str(self.tenant_id),
            "project_id": self.project_id,
        }

    def _normalize_loader_metadata(
        self,
        metadata: dict[str, Any],
        file_id: str,
        file_type: str,
    ) -> dict[str, Any]:
        normalized = self.clean_metadata(metadata)
        normalized.update(self._base_metadata(file_id, file_type))

        # PyMuPDFLoader returns a zero-based `page` value.
        raw_page = normalized.get("page")
        if isinstance(raw_page, int):
            normalized["page_index"] = raw_page
            normalized["page_number"] = raw_page + 1
        elif isinstance(raw_page, str) and raw_page.isdigit():
            page_index = int(raw_page)
            normalized["page_index"] = page_index
            normalized["page_number"] = page_index + 1

        return normalized

    def get_file_loader(self, file_id: str):
        file_extension = self.get_file_extension(file_id=file_id)
        file_path = self.get_file_path(file_id=file_id)
        if not os.path.isfile(file_path):
            return None

        if file_extension == ProcessingEnum.TXT.value:
            return TextLoader(
                file_path,
                encoding="utf-8",
                autodetect_encoding=True,
            )
        if file_extension == ProcessingEnum.PDF.value:
            return PyMuPDFLoader(file_path)
        return None

    def _load_docx(
        self,
        file_path: str,
        file_id: str,
    ) -> List[Document]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError(
                "DOCX support requires python-docx"
            ) from exc

        docx_document = DocxDocument(file_path)
        paragraphs = [
            paragraph.text.strip()
            for paragraph in docx_document.paragraphs
            if paragraph.text and paragraph.text.strip()
        ]
        text = "\n".join(paragraphs)
        if not text:
            return []

        metadata = self._base_metadata(file_id, "docx")
        metadata["source"] = file_path
        return [Document(page_content=text, metadata=metadata)]

    def _load_csv(
        self,
        file_path: str,
        file_id: str,
    ) -> List[Document]:
        documents: List[Document] = []
        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as source:
            reader = csv.DictReader(source)
            if reader.fieldnames:
                for row_number, row in enumerate(reader, start=2):
                    values = [
                        f"{key}: {value}"
                        for key, value in row.items()
                        if value is not None and str(value).strip()
                    ]
                    if not values:
                        continue
                    metadata = self._base_metadata(file_id, "csv")
                    metadata.update(
                        {
                            "source": file_path,
                            "row_start": row_number,
                            "row_end": row_number,
                        }
                    )
                    documents.append(
                        Document(
                            page_content="\n".join(values),
                            metadata=metadata,
                        )
                    )
            else:
                source.seek(0)
                plain_reader = csv.reader(source)
                for row_number, row in enumerate(
                    plain_reader,
                    start=1,
                ):
                    if not row:
                        continue
                    metadata = self._base_metadata(file_id, "csv")
                    metadata.update(
                        {
                            "source": file_path,
                            "row_start": row_number,
                            "row_end": row_number,
                        }
                    )
                    documents.append(
                        Document(
                            page_content=" | ".join(row),
                            metadata=metadata,
                        )
                    )
        return documents

    def _load_excel(
        self,
        file_path: str,
        file_id: str,
    ) -> List[Document]:
        try:
            import pandas as pd
        except ImportError as exc:
            raise RuntimeError(
                "Excel support requires pandas and openpyxl"
            ) from exc

        sheets = pd.read_excel(file_path, sheet_name=None)
        documents: List[Document] = []
        file_type = Path(file_path).suffix.lower().lstrip(".")

        for sheet_name, dataframe in sheets.items():
            dataframe = dataframe.fillna("")
            for index, row in dataframe.iterrows():
                values = [
                    f"{column}: {row[column]}"
                    for column in dataframe.columns
                    if str(row[column]).strip()
                ]
                if not values:
                    continue
                row_number = int(index) + 2
                metadata = self._base_metadata(file_id, file_type)
                metadata.update(
                    {
                        "source": file_path,
                        "sheet_name": str(sheet_name),
                        "row_start": row_number,
                        "row_end": row_number,
                    }
                )
                documents.append(
                    Document(
                        page_content="\n".join(values),
                        metadata=metadata,
                    )
                )
        return documents

    def get_file_content(
        self,
        file_id: str,
    ) -> List[Document] | None:
        file_extension = self.get_file_extension(file_id=file_id)
        file_path = self.get_file_path(file_id=file_id)
        if not os.path.isfile(file_path):
            return None

        try:
            if file_extension == ".docx":
                loaded_documents = self._load_docx(
                    file_path,
                    file_id,
                )
            elif file_extension == ".csv":
                loaded_documents = self._load_csv(
                    file_path,
                    file_id,
                )
            elif file_extension in {".xlsx", ".xls"}:
                loaded_documents = self._load_excel(
                    file_path,
                    file_id,
                )
            else:
                loader = self.get_file_loader(file_id=file_id)
                if loader is None:
                    return None
                loaded_documents = loader.load()
        except Exception as exc:
            raise RuntimeError(
                f"failed to load file: {file_id}"
            ) from exc

        cleaned_documents: List[Document] = []
        file_type = file_extension.lstrip(".")
        for record in loaded_documents:
            text = self.clean_text(
                getattr(record, "page_content", "")
            )
            if not text:
                continue
            metadata = self._normalize_loader_metadata(
                getattr(record, "metadata", {}),
                file_id=file_id,
                file_type=file_type,
            )
            cleaned_documents.append(
                Document(
                    page_content=text,
                    metadata=metadata,
                )
            )
        return cleaned_documents

    def process_file_content(
        self,
        file_content: List[Document],
        file_id: str,
        chunk_size: int = 500,
        overlap_size: int = 50,
    ) -> List[Document]:
        self._validate_chunk_settings(
            chunk_size=chunk_size,
            overlap_size=overlap_size,
        )
        if not file_content:
            return []

        texts: List[str] = []
        metadata_items: List[dict[str, Any]] = []
        for record in file_content:
            text = self.clean_text(
                getattr(record, "page_content", "")
            )
            if not text:
                continue
            metadata = self.clean_metadata(
                getattr(record, "metadata", {})
            )
            metadata.update(
                {
                    "file_id": file_id,
                    "stored_file_name": file_id,
                    "tenant_id": str(self.tenant_id),
                    "project_id": self.project_id,
                }
            )
            texts.append(text)
            metadata_items.append(metadata)

        if not texts:
            return []

        return self.process_simpler_splitter(
            texts=texts,
            metadatas=metadata_items,
            chunk_size=chunk_size,
            overlap_size=overlap_size,
        )

    def process_simpler_splitter(
        self,
        texts: List[str],
        metadatas: List[dict[str, Any]],
        chunk_size: int,
        overlap_size: int = 50,
    ) -> List[Document]:
        self._validate_chunk_settings(
            chunk_size=chunk_size,
            overlap_size=overlap_size,
        )
        chunks: List[Document] = []
        for index, raw_text in enumerate(texts):
            text = self.clean_text(raw_text)
            if not text:
                continue
            metadata = (
                metadatas[index]
                if index < len(metadatas)
                else {}
            )
            clean_metadata = self.clean_metadata(metadata)
            start = 0
            text_length = len(text)

            while start < text_length:
                hard_end = min(start + chunk_size, text_length)
                end = hard_end
                if hard_end < text_length:
                    candidate = text[start:hard_end]
                    break_position = max(
                        candidate.rfind("\n"),
                        candidate.rfind(" "),
                    )
                    minimum_break_position = int(chunk_size * 0.6)
                    if break_position >= minimum_break_position:
                        end = start + break_position

                chunk_text = self.clean_text(text[start:end])
                if chunk_text:
                    chunk_metadata = dict(clean_metadata)
                    chunk_metadata["chunk_start"] = start
                    chunk_metadata["chunk_end"] = end
                    chunks.append(
                        Document(
                            page_content=chunk_text,
                            metadata=chunk_metadata,
                        )
                    )

                if end >= text_length:
                    break
                next_start = end - overlap_size
                if next_start <= start:
                    next_start = end
                start = next_start
        return chunks

    @staticmethod
    def _validate_chunk_settings(
        chunk_size: int,
        overlap_size: int,
    ) -> None:
        if chunk_size <= 0:
            raise ValueError("chunk_size must be positive")
        if overlap_size < 0:
            raise ValueError("overlap_size cannot be negative")
        if overlap_size >= chunk_size:
            raise ValueError(
                "overlap_size must be smaller than chunk_size"
            )
